"""Card-format Word rendering: code cards, chips, adaptive options, PDF path.

Code stays real selectable text. Line numbers live in a separate gutter column
and never in the code runs. Syntax colors come from Pygments but never change
the underlying characters.
"""

import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from qbank_formats import FormatError

CHIP_SPAN = re.compile(r"`([^`]+)`")
LETTERS = "ABCD"


def shade_cell(cell, fill_hex):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill_hex)
    properties.append(shading)


def shade_run(run, fill_hex):
    run_properties = run._r.get_or_add_rPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill_hex)
    run_properties.append(shading)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("start", left),
                        ("end", right), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def remove_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def highlighted_lines(code_lines, language, highlighter):
    """Split lexer tokens into per-line runs; falls back to plain monospace."""
    try:
        from pygments import lex
        from pygments.lexers import get_lexer_by_name
        from pygments.styles import get_style_by_name
        lexer = get_lexer_by_name(language, stripall=False)
        style = get_style_by_name(highlighter)
    except Exception:
        return [[(None, line)] for line in code_lines], None
    code_text = "\n".join(code_lines)
    lines = [[]]
    for token_type, value in lex(code_text, lexer):
        parts = value.split("\n")
        for index, part in enumerate(parts):
            if index:
                lines.append([])
            if part:
                lines[-1].append((token_type, part))
    # Lexers emit a trailing newline token even when the text has none; drop
    # that phantom line while keeping a real trailing blank line.
    if not code_text.endswith("\n") and lines and not lines[-1]:
        lines.pop()
    return lines, style


def token_color(style, token_type, fallback):
    if style is None:
        return fallback
    color = style.style_for_token(token_type)["color"]
    return color or fallback


def add_code_runs(paragraph, line_runs, tokens, style):
    for token_type, value in line_runs:
        run = paragraph.add_run(value)
        run.font.name = tokens["font"]
        run.font.size = Pt(tokens["font_size_pt"])
        run.font.color.rgb = RGBColor.from_string(token_color(style, token_type, tokens["code_text"]))
        if style is not None:
            token_style = style.style_for_token(token_type)
            run.font.bold = bool(token_style.get("bold"))
            run.font.italic = bool(token_style.get("italic"))


def gutter_width_in(lines_count):
    digits = max(len(str(lines_count)), 1)
    return 0.26 + 0.08 * max(digits - 1, 0)


def keep_row_intact(row):
    """Prevent a table row from splitting across pages (Word ignores this when
    the row is taller than a page, so long listings still split safely)."""
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_code_card(document, code_lines, language, header_left, header_right, tokens,
                  highlighter):
    """One card as a single table: header row, then gutter and code columns.

    A single table keeps Word from merging adjacent tables and keeps the header
    attached to its code. Code characters stay real, selectable text.
    """
    lines, style = highlighted_lines(code_lines, language, highlighter)
    table = document.add_table(rows=2, cols=2)
    remove_table_borders(table)
    table.autofit = False
    keep_row_intact(table.rows[0])
    keep_row_intact(table.rows[1])
    gutter, code = table.rows[1].cells
    total_width = 6.1
    gutter_in = gutter_width_in(len(lines))
    gutter.width = Inches(gutter_in)
    code.width = Inches(total_width - gutter_in)
    # The fixed layout resolves an ambiguous grid (default equal columns) to
    # different widths per table, overflowing the margin; align gridCol with
    # the designed cell widths so every card lays out identically.
    table.columns[0].width = Inches(gutter_in)
    table.columns[1].width = Inches(total_width - gutter_in)
    header_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    header_cell.width = Inches(total_width)
    shade_cell(header_cell, tokens["header"])
    shade_cell(gutter, tokens["gutter"])
    shade_cell(code, tokens["code"])
    set_cell_margins(gutter, left=20, right=30)
    set_cell_margins(code, left=50)
    gutter.vertical_alignment = code.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    header_paragraph = header_cell.paragraphs[0]
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_paragraph.paragraph_format.keep_with_next = True
    header_paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(total_width - 0.15), WD_TAB_ALIGNMENT.RIGHT)
    run = header_paragraph.add_run(header_left)
    run.font.size = Pt(9)
    run.bold = True
    run = header_paragraph.add_run("\t" + header_right)
    run.font.size = Pt(9)
    run.bold = True

    for index in range(len(lines)):
        number_paragraph = gutter.paragraphs[0] if index == 0 else gutter.add_paragraph()
        run = number_paragraph.add_run(str(index + 1))
        run.font.name = tokens["font"]
        run.font.size = Pt(tokens["font_size_pt"])
        run.font.color.rgb = RGBColor.from_string(tokens["gutter_text"])
        number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number_paragraph.paragraph_format.space_after = Pt(0)
        code_paragraph = code.paragraphs[0] if index == 0 else code.add_paragraph()
        add_code_runs(code_paragraph, lines[index], tokens, style)
        code_paragraph.paragraph_format.space_after = Pt(0)
    # Keep the card attached to the question that follows it: a card must never
    # sit orphaned at a page bottom with its questions on the next page.
    code_paragraph.paragraph_format.keep_with_next = True
    return table


def add_chip_text(paragraph, text, tokens, base_size_pt=11, bold=False):
    """Render text with backtick code spans as shaded monospace chips."""
    position = 0
    for match in CHIP_SPAN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.font.size = Pt(base_size_pt)
            run.bold = bold
        chip = paragraph.add_run(match.group(1))
        chip.font.name = tokens["font"]
        chip.font.size = Pt(tokens["chip_font_size_pt"])
        chip.font.color.rgb = RGBColor.from_string("24292f")
        shade_run(chip, tokens["chip"])
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.font.size = Pt(base_size_pt)
        run.bold = bold


def visible_length(text):
    return len(CHIP_SPAN.sub(lambda match: match.group(1), text))


def add_options_block(document, options, tokens):
    """MCQ options: two by two when short, one per row when long."""
    styled = [(letter, text.strip()) for letter, text in zip(LETTERS, options)]
    threshold = tokens["option_short_threshold"]
    if tokens["option_layout"] == "adaptive" and len(styled) == 4 and all(
            visible_length(text) <= threshold for _, text in styled):
        table = document.add_table(rows=2, cols=2)
        remove_table_borders(table)
        table.autofit = False
        for index, (letter, text) in enumerate(styled):
            cell = table.rows[index // 2].cells[index % 2]
            paragraph = cell.paragraphs[0]
            add_chip_text(paragraph, f"{letter}) {text}", tokens)
            paragraph.paragraph_format.space_after = Pt(2)
        # A question's options never split across pages: rows cannot split and the
        # first row stays with the second.
        for row in table.rows:
            keep_row_intact(row)
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        return table
    for index, (letter, text) in enumerate(styled):
        paragraph = document.add_paragraph(style="List Bullet 2")
        paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
        add_chip_text(paragraph, f"{letter}) {text}", tokens)
        paragraph.paragraph_format.keep_with_next = index < len(styled) - 1
    return None


def convert_docx_to_pdf(docx_path, pdf_path):
    """Convert the finalized DOCX to PDF with a local renderer; honest result.

    Returns {"status": "converted", "renderer": ...} or
    {"status": "unavailable", "reason": ...}. A separately drawn PDF is never
    substituted for the converted document.
    """
    from pathlib import Path

    docx_path, pdf_path = Path(docx_path), Path(pdf_path)
    try:
        return convert_with_word(docx_path, pdf_path)
    except ImportError:
        word_error = "pywin32 is not installed in this Python environment"
    except Exception as error:  # COM errors are environment-specific disclosures
        word_error = f"Microsoft Word COM conversion failed: {error}"
    fallback = convert_with_soffice(docx_path, pdf_path)
    if fallback is not None:
        return fallback
    return {"status": "unavailable", "reason": word_error or
            "No local Word/LibreOffice converter is available"}


def convert_with_word(docx_path, pdf_path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        document = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        document.ExportAsFixedFormat(OutputFileName=str(pdf_path.resolve()),
                                     ExportFormat=17, OpenAfterExport=False,
                                     OptimizeFor=0, CreateBookmarks=0)
        document.Close(SaveChanges=0)
        version = word.Version
    finally:
        word.Quit()
        pythoncom.CoUninitialize()
    return verify_pdf(pdf_path, f"Microsoft Word {version} (COM ExportAsFixedFormat)")


def convert_with_soffice(docx_path, pdf_path):
    import shutil
    import subprocess

    soffice = shutil.which("soffice")
    if soffice is None:
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(pdf_path.parent), str(docx_path)], check=True, timeout=180)
    return verify_pdf(pdf_path, f"LibreOffice ({soffice})")


def verify_pdf(pdf_path, renderer):
    from pypdf import PdfReader

    if not pdf_path.exists() or pdf_path.stat().st_size == 0:
        raise FormatError(f"{renderer} produced no PDF at {pdf_path}")
    pages = len(PdfReader(str(pdf_path)).pages)
    if pages < 1:
        raise FormatError(f"{renderer} produced an empty PDF")
    return {"status": "converted", "renderer": renderer, "pages": pages,
            "pdf_path": str(pdf_path)}
