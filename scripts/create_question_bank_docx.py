import argparse
import copy
import json
import random
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = SKILL_DIR / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from qbank_document import DEFAULT_TEMPLATE, build_pdf_from_docx, build_v2_docx  # noqa: E402
from qbank_schema import (SOURCE_MODES, SchemaError, normalized_text,  # noqa: E402
                          validate_legacy_bank)

# Shared validation lives in qbank_schema; keep the historical name working.
validate_bank = validate_legacy_bank

DIFFICULTIES = {"easy", "medium", "hard"}
SECTION_INSTRUCTIONS = {
    "mcqs": "Choose one answer for each question.",
    "true_false": "Mark each statement True or False.",
    "essays": "Answer every part of each prompt in your own words."
}
SECTION_TITLES = {
    "en": {"mcqs": "Multiple Choice Questions", "true_false": "True or False Questions",
           "essays": "Essay Questions", "headers": ["No.", "Answer", "Explanation or correction", "Source"]}
}


def shuffled_sections(bank):
    sections = copy.deepcopy({section: bank.get(section, []) for section in
                              ("mcqs", "true_false", "essays")})
    rng = random.SystemRandom()
    for section, questions in sections.items():
        rng.shuffle(questions)
        if section == "mcqs":
            for question in questions:
                correct_text = question["options"]["ABCD".index(question["correct"])]
                rng.shuffle(question["options"])
                question["correct"] = "ABCD"[question["options"].index(correct_text)]
    return sections


def set_text_direction(paragraph):
    if re.search(r"[\u0600-\u06ff]", paragraph.text):
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def add_question(document, number, stem):
    paragraph = document.add_paragraph(style="List Number")
    # Explicit numbering keeps section headings and answer rows in agreement.
    num_properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_properties.get_or_add_numId().val = 0
    paragraph.add_run(f"{number}. {stem.strip()}").bold = True
    paragraph.paragraph_format.keep_with_next = True
    set_text_direction(paragraph)


def add_mcq_body(document, question):
    for letter, option in zip("ABCD", question["options"]):
        paragraph = document.add_paragraph(f"{letter}) {option.strip()}", style="List Bullet 2")
        paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
        paragraph.paragraph_format.keep_with_next = letter != "D"
        set_text_direction(paragraph)


def add_true_false_body(document, question):
    document.add_paragraph("True / False")


def add_essay_body(document, question):
    document.add_paragraph("................................................................................")


QUESTION_RENDERERS = {"mcqs": add_mcq_body, "true_false": add_true_false_body,
                      "essays": add_essay_body}


def question_answer(question, section):
    if section == "mcqs":
        letter = question["correct"]
        return f"{letter}) {question['options']['ABCD'.index(letter)].strip()}"
    if section == "true_false":
        return question["correct"]
    return "\n".join(f"• {point.strip()}" for point in question["model_answer"])


def source_key(source):
    return source["document"], source["kind"].removesuffix("_inference")


def source_registry(sections):
    registry = {}
    for questions in sections.values():
        for question in questions:
            for source in question["sources"]:
                if source["kind"] not in {"web", "general"}:
                    key = source_key(source)
                    if key not in registry:
                        registry[key] = f"S{len(registry) + 1}"
    return registry


def source_label(source, registry=None, source_mode="pdf_only"):
    kind = source["kind"]
    if kind not in {"web", "general"}:
        identifier = registry[source_key(source)] if registry else source["document"]
        base_kind = kind.removesuffix("_inference")
        label = f"{identifier}: "
        if base_kind == "pdf":
            label += f"PDF page {source['pdf_page']}"
        elif base_kind == "slides":
            label += f"slide {source['slide']}"
        else:
            if source.get("page"):
                label += f"page {source['page']}, "
            label += source["locator"]
        if source.get("printed_page"):
            label += f", printed page {source['printed_page']}"
        if base_kind != "document" and source.get("locator"):
            label += f", {source['locator']}"
        if kind.endswith("_inference"):
            origin = {"pdf": "PDF", "slides": "slides", "document": "source"}[base_kind]
            label += f"\nInference from {origin}"
        return label
    outside = "Outside PDF" if source_mode.startswith("pdf_") else "Outside uploaded material"
    if kind == "web":
        return f"{outside}: Web: {source['title']}\n{source['url']}"
    return f"{outside}: General knowledge from model; checked against {source['verified_against']}"


def answer_row(question, section, number, registry=None, source_mode="pdf_only"):
    explanation = question["explanation"].strip()
    if section == "true_false" and question["correct"] == "False":
        explanation = "Correction: " + question["correction"].strip()
    labels = list(dict.fromkeys(source_label(source, registry, source_mode) for source in question["sources"]))
    return [str(number), question_answer(question, section), explanation,
            "\n".join(labels)]


def add_question_sections(document, sections, language, registry=None, source_mode="pdf_only"):
    answers = []
    for section, questions in sections.items():
        if not questions:
            continue
        heading = document.add_paragraph(SECTION_TITLES[language][section], style="Heading 1")
        if answers:
            heading.paragraph_format.page_break_before = True
        set_text_direction(heading)
        instruction = document.add_paragraph(SECTION_INSTRUCTIONS[section])
        instruction.paragraph_format.keep_with_next = True
        for question in questions:
            number = len(answers) + 1
            stem = question.get("question", question.get("statement"))
            add_question(document, number, stem)
            QUESTION_RENDERERS[section](document, question)
            answers.append(answer_row(question, section, number, registry, source_mode))
    return answers


def fill_answer_cells(cells, content):
    for cell, text in zip(cells, content):
        cell.text = text
        for paragraph in cell.paragraphs:
            set_text_direction(paragraph)
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(10)


def add_answer_key(document, answers, language, registry=None):
    document.add_page_break()
    document.add_paragraph("Model Answer", style="Heading 1")
    if registry:
        document.add_paragraph("Source references below identify the uploaded files. PDF pages and slides use 1-based numbering. Inference labels identify applications or deductions from the cited material.")
        for (name, _), identifier in registry.items():
            document.add_paragraph(f"{identifier}: {name}")
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    table.autofit = False
    for column, width in zip(table.columns, (0.45, 1.65, 1.8, 2.1)):
        column.width = Inches(width)
    header = table.rows[0]
    for cell, column in zip(header.cells, table.columns):
        cell.width = column.width
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    fill_answer_cells(header.cells, SECTION_TITLES[language]["headers"])
    for cell in header.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for answer in answers:
        fill_answer_cells(table.add_row().cells, answer)


def build_docx(bank, output_path, template_path=DEFAULT_TEMPLATE,
               format_id=None, variant=None):
    """Export an audited bank. Version 2 banks dispatch to the new pipeline."""
    if isinstance(bank, dict) and bank.get("schema_version") == 2:
        return build_v2_docx(bank, output_path, template_path, format_id, variant)
    validate_legacy_bank(bank)
    document = Document(str(template_path))
    for child in list(document._body._element):
        if child.tag != qn("w:sectPr"):
            document._body._element.remove(child)
    title = document.add_paragraph(bank["title"].strip(), style="Title")
    set_text_direction(title)
    language = bank.get("language", "en")
    document.add_paragraph(f"{sum(bank['counts'].values())} questions: {bank['counts']['mcqs']} multiple choice, {bank['counts']['true_false']} True/False, and {bank['counts']['essays']} essay.")
    sections = shuffled_sections(bank)
    registry = source_registry(sections)
    answers = add_question_sections(document, sections, language, registry, bank["source_mode"])
    if bank.get("include_answer_key", True):
        add_answer_key(document, answers, language, registry)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def main():
    parser = argparse.ArgumentParser(description="Export an audited document or slide question bank as DOCX (and optionally PDF).")
    parser.add_argument("input_json")
    parser.add_argument("output_docx")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--format", default=None, dest="format_id",
                        help="Document format for version-2 banks: classic (default) or a Card/saved format id")
    parser.add_argument("--variant", default=None,
                        help="Card variant: monokai (default) or print")
    parser.add_argument("--pdf", default=None, metavar="OUTPUT_PDF",
                        help="Also convert the finalized DOCX to this PDF path")
    arguments = parser.parse_args()
    with open(arguments.input_json, encoding="utf-8") as question_file:
        bank = json.load(question_file)
    build_docx(bank, arguments.output_docx, Path(arguments.template),
               arguments.format_id, arguments.variant)
    if arguments.pdf:
        result = build_pdf_from_docx(arguments.output_docx, arguments.pdf)
        if result["status"] != "converted":
            parser.exit(2, f"PDF conversion unavailable: {result['reason']}\n")
        print(result)


if __name__ == "__main__":
    main()
