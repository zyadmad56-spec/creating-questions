"""Extract study material without rendering pages or uploading images.

Text-first extraction with code recovery: detect code candidates, record the
language evidence, preserve original locators and text, and flag review states
instead of silently rewriting source meaning.
"""

import argparse
import ast
import json
import re
from pathlib import Path


CODE_LANGUAGE_LABELS = {
    "python": "python", "c": "c", "cpp": "cpp", "java": "java",
    "csharp": "csharp", "javascript": "javascript", "sql": "sql",
    "r": "r", "html": "html",
}

# Ordered strongest-first: a match records evidence for the language.
LANGUAGE_INDICATORS = {
    "python": [r"\bdef\s+\w+\s*\(", r"\bclass\s+\w+\s*[:(]", r"\bimport\s+\w+",
               r"\bfrom\s+\w+\s+import\b", r"\bprint\s*\(", r"\bself\.", r"\belif\b",
               r"\bNone\b", r"\bTrue\b", r"\borchestrat\b"],
    "c": [r"#include\s*<\w+\.h>", r"\bprintf\s*\(", r"\bscanf\s*\(",
          r"\bmalloc\s*\(", r"\bstruct\s+\w+\s*\{"],
    "cpp": [r"#include\s*<iostream>", r"\bstd::", r"\bcout\s*<<", r"\bcin\s*>>",
            r"\bnamespace\s+\w+", r"\btemplate\s*<"],
    "java": [r"\bpublic\s+class\b", r"System\.out\.print", r"\bpublic\s+static\s+void\s+main\b",
             r"\bprivate\s+\w+\s+\w+;", r"\bextends\b"],
    "csharp": [r"\busing\s+System\b", r"Console\.Write", r"\bnamespace\s+\w+",
               r"\bpublic\s+void\b"],
    "javascript": [r"\bfunction\s+\w+\s*\(", r"\bconst\s+\w+\s*=", r"\blet\s+\w+\s*=",
                   r"console\.log", r"=>", r"\bdocument\."],
    "sql": [r"\bSELECT\b.+\bFROM\b", r"\bINSERT\s+INTO\b", r"\bCREATE\s+TABLE\b",
            r"\bUPDATE\s+\w+\s+SET\b", r"\bDELETE\s+FROM\b", r"\bJOIN\b"],
    "r": [r"<-\s*", r"\blibrary\s*\(", r"\bggplot\b", r"\bdata\.frame\b"],
    "html": [r"<html\b", r"<div\b", r"</p>", r"<body\b"],
}

VERSION_HINT_PATTERNS = [
    (r"python\s*3(?:\.\d+)?", "3.x"),
    (r"C\+\+\s*(?:1[147]|20)", None),
    (r"\bC\s*1[179]9?\b", None),
    (r"\bJava\s*(?:8|1[1-9]|2[0-9])\b", None),
    (r"MySQL\s*\d", None),
    (r"PostgreSQL\s*\d+", None),
]

CONSTRUCT_PATTERNS = {
    "python": [
        ("functions", r"\bdef\s+\w+\s*\("), ("classes", r"\bclass\s+\w+"),
        ("loops", r"\bfor\s+\w+\s+in\b|\bwhile\s+.+:"),
        ("conditionals", r"\bif\s+.+:"),
        ("exception_handling", r"\btry\b:|\bexcept\b|\braise\b"),
        ("imports", r"^\s*import\s+|^\s*from\s+\w+\s+import", ),
        ("list_comprehensions", r"\[\s*\w+\s+for\s+\w+\s+in"),
        ("methods", r"\bdef\s+self|def\s+__init__"),
        ("attribute_access", r"\bself\.\w+"),
    ],
    "c": [("functions", r"\w+\s+\w+\s*\([^;]*\)\s*\{"), ("pointers", r"\*\w+|\b&\w+"),
          ("loops", r"\bfor\s*\(|\bwhile\s*\("), ("conditionals", r"\bif\s*\("),
          ("structs", r"\bstruct\s+\w+"), ("printf_output", r"\bprintf\s*\(")],
    "cpp": [("classes", r"\bclass\s+\w+"), ("iostream_io", r"\bcout\b|\bcin\b"),
            ("loops", r"\bfor\s*\(|\bwhile\s*\("), ("conditionals", r"\bif\s*\("),
            ("stl_containers", r"\bstd::(vector|map|string|set)"),
            ("templates", r"\btemplate\s*<")],
    "java": [("classes", r"\bclass\s+\w+"), ("methods", r"\b(public|private|protected)\s+\w+.*\("),
             ("loops", r"\bfor\s*\(|\bwhile\s*\("), ("conditionals", r"\bif\s*\("),
             ("inheritance", r"\bextends\b|\bimplements\b"),
             ("console_output", r"System\.out\.print")],
    "csharp": [("classes", r"\bclass\s+\w+"), ("methods", r"\b(public|private)\s+\w+.*\("),
               ("loops", r"\bfor\s*\(|\bwhile\s*\("), ("console_output", r"Console\.Write")],
    "javascript": [("functions", r"\bfunction\b|=>"), ("objects", r"\{\s*\w+\s*:"),
                   ("loops", r"\bfor\s*\(|\bwhile\s*\(|\bforEach\b"),
                   ("dom_access", r"\bdocument\."), ("async", r"\bawait\b|\basync\b")],
    "sql": [("select_queries", r"\bSELECT\b"), ("joins", r"\bJOIN\b"),
            ("aggregation", r"\bGROUP\s+BY\b|\bCOUNT\s*\(|\bSUM\s*\("),
            ("data_definition", r"\bCREATE\s+(TABLE|INDEX|VIEW)\b"),
            ("data_modification", r"\bINSERT\s+INTO\b|\bUPDATE\b|\bDELETE\s+FROM\b")],
}

STRONG_CODE_MARKERS = re.compile(
    r"\b(def|class|import|return|elif|function|const|let|public\s+class|System\.out|"
    r"printf|cout|#include|SELECT|FROM|INSERT|console\.log)\b|\{$|;\s*$|=>")
WEAK_CODE_MARKERS = re.compile(
    r"^\s{2,}\S|^\t|^\s*[A-Za-z_]\w*\s*=[^=]|[A-Za-z_]\w*(\.\w+)*\s*=[^=]|"
    r"\)\s*:?\s*$|\)\s*;?\s*$|\bfor\b|\bif\b|\bwhile\b|\breturn\b|\bprint\b|"
    r"^\s*\}\s*$|^\s*\{\s*$|\w\s*(\+|-|\*|/)=(?!=)|\w\s*<-\s*\w")
GAP_MARKER = re.compile(r"_{3,}|\. {2,}\.|\bTODO\b|\bpass\s*$")
# One separating space after the label is display formatting; further spaces are
# code indentation and must survive into the normalized text.
DISPLAY_LABEL = re.compile(r"^(\s*)(\d{1,3})\s*[.\-):]\s")


def markdown_table(rows):
    if not rows:
        return ""
    cells = [[str(value).replace("|", "\\|").replace("\n", "<br>") for value in row] for row in rows]
    lines = ["| " + " | ".join(cells[0]) + " |",
             "| " + " | ".join("---" for _ in cells[0]) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in cells[1:])
    return "\n".join(lines)


def detect_language(code_text, context_text=""):
    """Return language, confidence, evidence tokens, and a version hint if stated."""
    scores = {}
    for language, patterns in LANGUAGE_INDICATORS.items():
        evidence = [pattern for pattern in patterns
                    if re.search(pattern, code_text, re.MULTILINE)]
        if evidence:
            scores[language] = evidence
    if not scores:
        return {"language": None, "confidence": 0, "evidence": [], "version_hint": None}
    language = max(scores, key=lambda name: (len(scores[name]), name))
    hint = None
    for pattern, fixed in VERSION_HINT_PATTERNS:
        match = re.search(pattern, context_text, re.IGNORECASE)
        if match:
            hint = fixed or match.group(0)
            break
    return {"language": language, "confidence": len(scores[language]),
            "evidence": scores[language][:4], "version_hint": hint}


PROSE_SENTENCE = re.compile(r"[.!?]$")
CODE_PUNCTUATION = re.compile(r"[{}<>=;:]|\(|\)|\[")


def line_is_code_like(line):
    """Judge code likeness on the text a reader sees, without printed labels,
    keeping the code's own indentation intact."""
    remainder = DISPLAY_LABEL.sub(lambda match: match.group(1), line)
    stripped = remainder.strip()
    if not stripped:
        return True
    if PROSE_SENTENCE.search(stripped) and not CODE_PUNCTUATION.search(stripped):
        return False
    return bool(STRONG_CODE_MARKERS.search(stripped) or WEAK_CODE_MARKERS.search(remainder))


def block_is_code(lines):
    """A run is code when strong markers or consistent indentation dominate it."""
    nonempty = [line for line in lines if line.strip()]
    if not nonempty:
        return False
    strong = sum(1 for line in nonempty if STRONG_CODE_MARKERS.search(line.strip()))
    weak = sum(1 for line in nonempty if line_is_code_like(line))
    indented = sum(1 for line in nonempty if line[:1] in (" ", "\t"))
    return strong >= 1 and weak >= len(nonempty) * 0.6 or indented >= len(nonempty) * 0.8 and weak >= 2


def split_labeled_lines(lines):
    """Remove printed display labels (01., 2-, 3)) and map them to line indexes."""
    normalized, labels, edits = [], [], []
    for index, line in enumerate(lines):
        match = DISPLAY_LABEL.match(line)
        if match:
            labels.append({"display_label": int(match.group(2)), "line_index": index})
            line = match.group(1) + line[match.end():]
        normalized.append(line.rstrip())
    if labels:
        edits.append({"edit": "removed printed display-number prefixes",
                      "reason": "Display labels are not executable text; mapping is recorded.",
                      "labels": labels})
    return normalized, labels, edits


def candidate_from_lines(unit, lines, first_line_number=None, source_kind="text"):
    normalized, labels, edits = split_labeled_lines(lines)
    original = "\n".join(lines)
    detection = detect_language("\n".join(normalized), unit.get("context", ""))
    tabbed = any("\t" in line for line in lines)
    if tabbed:
        edits.append({"edit": "preserved tab characters as recovered",
                      "reason": "Indentation is source evidence; conversion is a later, recorded decision."})
    if GAP_MARKER.search("\n".join(normalized)):
        review = "needs_review"
        edits.append({"edit": "flagged a blank/gap marker",
                      "reason": "The snippet may be an intentionally incomplete exercise; nothing was filled."})
    elif detection["language"] is None:
        review = "needs_review"
    else:
        review = "recovered"
    start = first_line_number + 1 if first_line_number is not None else None
    end = start + len(lines) - 1 if start is not None else None
    return {
        "locator": unit["locator"],
        "original_text": original,
        "normalized_code": "\n".join(normalized),
        "language": detection["language"],
        "version_hint": detection["version_hint"],
        "language_confidence": detection["confidence"],
        "language_evidence": detection["evidence"],
        "line_labels": labels,
        "recovery_edits": edits,
        "review_status": review,
        "source_lines": f"lines {start}-{end}" if start is not None else unit["locator"],
        "source_kind": source_kind,
    }


def extract_fenced_candidates(unit):
    candidates = []
    fence = re.compile(r"^```(\w*)\s*$", re.MULTILINE)
    matches = list(fence.finditer(unit["body"]))
    for index in range(0, len(matches) - 1, 2):
        opening, closing = matches[index], matches[index + 1]
        code = unit["body"][opening.end():closing.start()].rstrip("\n")
        lines = code.splitlines()
        record = candidate_from_lines(unit, lines)
        stated = opening.group(1).lower()
        record["stated_language"] = stated or None
        if stated in CODE_LANGUAGE_LABELS:
            record["language"] = stated
            record["language_confidence"] = max(record["language_confidence"], 2)
            record["language_evidence"] = [f"fenced block labeled {stated}"]
        candidates.append(record)
    return candidates


def scan_unit_for_candidates(unit, min_run=3):
    """Detect code-like line runs inside one extraction unit; blank lines inside
    an active run stay part of it so code blocks with empty lines survive."""
    candidates = []
    lines = unit["body"].splitlines()
    run_start, run_end, first_line_number = None, None, None
    for index, line in enumerate(lines):
        if line.strip() and line_is_code_like(line):
            if run_start is None:
                run_start, first_line_number = index, index
            run_end = index
        elif run_start is not None and not line.strip():
            continue  # blank line inside an active run
        elif run_start is not None:
            run = lines[run_start:run_end + 1]
            if len(run) >= min_run or block_is_code(run):
                candidates.append(candidate_from_lines(unit, run, first_line_number))
            run_start, run_end, first_line_number = None, None, None
    if run_start is not None:
        run = lines[run_start:run_end + 1]
        if len(run) >= min_run or block_is_code(run):
            candidates.append(candidate_from_lines(unit, run, first_line_number))
    return candidates


def group_consecutive_units(units, code_line_cache):
    """Merge consecutive code-like DOCX paragraphs into one candidate."""
    merged, run = [], []
    for unit in units:
        if code_line_cache.get(id(unit)):
            run.append(unit)
        elif run:
            merged.append(run)
            run = []
        else:
            merged.append([unit])
    if run:
        merged.append(run)
    return merged


def scan_units_for_candidates(units, merge_units=False):
    candidates = []
    unit_is_code_like = {}
    for unit in units:
        local = scan_unit_for_candidates(unit)
        unit_is_code_like[id(unit)] = merge_units and any(
            line.strip() and line_is_code_like(line) for line in unit["body"].splitlines())
        candidates.extend(local)
    if merge_units:
        candidates = []
        for group in group_consecutive_units(units, unit_is_code_like):
            if len(group) == 1:
                candidates.extend(scan_unit_for_candidates(group[0]))
            else:
                lines, locator = [], group[0]["locator"]
                for unit in group:
                    lines.extend(unit["body"].splitlines())
                candidates.append(candidate_from_lines(
                    {"locator": locator,
                     "context": "\n".join(unit["context"] for unit in group)},
                    lines, source_kind="consecutive_paragraphs"))
    return candidates


def check_python_syntax(candidate):
    if candidate["language"] != "python":
        return None
    try:
        ast.parse(candidate["normalized_code"])
    except SyntaxError as error:
        return {"checker": "ast.parse", "result": "syntax_error",
                "detail": f"line {error.lineno}: {error.msg}",
                "note": ("May be intentional (fill-the-gap exercise, pseudocode, partial method, "
                         "or deliberately broken program). Review state is recorded; nothing was rewritten.")}
    return {"checker": "ast.parse", "result": "passed", "detail": None, "note": None}


def inventory_constructs(language, code_texts, context_text):
    patterns = CONSTRUCT_PATTERNS.get(language, [])
    inventory = []
    for name, pattern in patterns:
        if re.search(pattern, "\n".join(code_texts), re.MULTILINE):
            inventory.append(name)
    if language not in CONSTRUCT_PATTERNS:
        inventory = None
    return inventory


def recover_code(units, merge_units=False):
    """Detect code candidates across units; returns (sections, report lists)."""
    candidates, fences = [], []
    for unit in units:
        fences.extend(extract_fenced_candidates(unit))
    candidates.extend(fences)
    plain = scan_units_for_candidates(units, merge_units=merge_units)
    # Keep fenced blocks first: they are explicit; plain scans fill the rest.
    seen_spans = {(fence["locator"], fence["normalized_code"]) for fence in fences}
    for candidate in plain:
        key = (candidate["locator"], candidate["normalized_code"])
        if key not in seen_spans:
            candidates.append(candidate)

    syntax_checks, languages = [], {}
    for number, candidate in enumerate(candidates, 1):
        candidate["candidate_id"] = f"cand_{number:03d}"
        check = check_python_syntax(candidate)
        if check:
            candidate["syntax_check"] = check["result"]
            syntax_checks.append({"candidate_id": candidate["candidate_id"], **check})
            if check["result"] == "syntax_error" and candidate["review_status"] == "recovered":
                candidate["review_status"] = "needs_review"
        inventory = inventory_constructs(candidate["language"], [candidate["normalized_code"]], "")
        if candidate["language"]:
            entry = languages.setdefault(candidate["language"], {
                "candidate_ids": [], "constructs": [], "confidence": 0})
            entry["candidate_ids"].append(candidate["candidate_id"])
            entry["confidence"] = max(entry["confidence"], candidate["language_confidence"])
            for construct in inventory or []:
                if construct not in entry["constructs"]:
                    entry["constructs"].append(construct)

    recovery_edits = [{"candidate_id": candidate["candidate_id"], **edit}
                      for candidate in candidates for edit in candidate["recovery_edits"]]
    sections = {}
    for candidate in candidates:
        sections.setdefault(candidate["locator"], []).append(candidate)
    return sections, candidates, languages, syntax_checks, recovery_edits


def render_candidate_section(candidate):
    label = candidate["language"] or ""
    lines = [
        f"### Recovered code candidate {candidate['candidate_id']} "
        f"({candidate['source_lines']}; language: {label or 'undetected'}; "
        f"review: {candidate['review_status']})",
        "",
        f"```{label}",
        candidate["normalized_code"],
        "```",
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Format extractors: each returns (units, review, unit_count)
# ---------------------------------------------------------------------------

def extract_pdf(path, layout_pages=None):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("The PDF is encrypted. Provide an accessible copy.")
    units, review = [], []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        body = text or "[No extractable text on this page.]"
        units.append({"locator": f"PDF page {number}", "heading": f"## PDF page {number}",
                      "body": body, "context": body})
        if layout_pages and number in layout_pages:
            layout_text = (page.extract_text(extraction_mode="layout") or "").strip()
            units.append({"locator": f"PDF page {number} (layout extraction)",
                          "heading": f"## PDF page {number} (layout extraction)",
                          "body": layout_text or "[Layout extraction produced no text.]",
                          "context": body})
        if len(text.split()) < 15:
            review.append({"pdf_page": number, "reason": "Little or no extracted text; check relevance before OCR or a targeted page view."})
    return units, review, len(reader.pages)


def extract_slide_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    parts, review = [], []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            nested, notices = extract_slide_shapes(shape.shapes)
            parts.extend(nested)
            review.extend(notices)
        elif shape.has_table:
            parts.append(markdown_table([[cell.text for cell in row.cells] for row in shape.table.rows]))
        elif shape.has_text_frame:
            text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
            if text:
                parts.append(text)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            review.append(f"Picture present: {shape.name}; its contents were not interpreted.")
        elif getattr(shape, "has_chart", False):
            review.append(f"Chart present: {shape.name}; values need a source check if used.")
    return parts, review


def extract_pptx(path):
    from pptx import Presentation

    deck = Presentation(path)
    units, review = [], []
    for number, slide in enumerate(deck.slides, 1):
        content, notices = extract_slide_shapes(slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                content.append("### Speaker notes\n\n" + notes)
        if not content:
            content.append("[No extractable text or tables on this slide.]")
            notices.append("No text or tables extracted; check this slide if it belongs to the selected scope.")
        body = "\n\n".join(content)
        units.append({"locator": f"Slide {number}", "heading": f"## Slide {number}",
                      "body": body, "context": body})
        if notices:
            review.append({"slide": number, "reasons": notices})
    return units, review, len(deck.slides)


def extract_docx(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    units, review = [], []
    paragraph_number = table_number = 0
    for node in document.element.body:
        if node.tag == qn("w:p"):
            paragraph_number += 1
            paragraph = Paragraph(node, document)
            if paragraph.text.strip():
                units.append({"locator": f"Paragraph {paragraph_number}",
                              "heading": f"## Paragraph {paragraph_number}",
                              "body": paragraph.text, "context": paragraph.text})
            if node.findall(".//" + qn("w:drawing")) or node.findall(".//" + qn("w:pict")):
                review.append({"locator": f"Paragraph {paragraph_number}",
                               "reason": "Drawing or image present; inspect only if it contains needed evidence."})
        elif node.tag == qn("w:tbl"):
            table_number += 1
            table = Table(node, document)
            units.append({"locator": f"Table {table_number}",
                          "heading": f"## Table {table_number}",
                          "body": markdown_table([[cell.text for cell in row.cells] for row in table.rows]),
                          "context": ""})
            if node.findall(".//" + qn("w:drawing")) or node.findall(".//" + qn("w:pict")):
                review.append({"locator": f"Table {table_number}", "reason": "Drawing or image present in this table."})
    return units, review, paragraph_number + table_number


def extract_text(path):
    text = path.read_text(encoding="utf-8-sig")
    return ([{"locator": "original document text", "heading": "## Original document text",
              "body": text, "context": ""}], [], len(text.splitlines()))


EXTRACTORS = {".pdf": extract_pdf, ".pptx": extract_pptx, ".docx": extract_docx,
              ".md": extract_text, ".txt": extract_text}

MERGED_UNIT_FORMATS = {".docx"}


def parse_layout_pages(layout_mode, units):
    """Resolve --layout into an explicit page set: off, auto, or listed numbers."""
    if layout_mode == "off" or not layout_mode:
        return None
    if layout_mode != "auto":
        try:
            return {int(item) for item in layout_mode.split(",") if item.strip()}
        except ValueError as error:
            raise ValueError("--layout accepts off, auto, or comma-separated page numbers") from error
    flagged = set()
    for unit in units:
        match = re.match(r"PDF page (\d+)$", unit["locator"])
        if match and scan_unit_for_candidates(unit):
            flagged.add(int(match.group(1)))
    return flagged


def extract_source(source_path, output_path, layout_mode="off"):
    source, output = Path(source_path), Path(output_path)
    if source.resolve() == output.resolve():
        raise ValueError("The extraction must not overwrite its source.")
    if source.suffix.lower() == ".ppt":
        raise ValueError("Legacy .ppt needs conversion to .pptx or PDF using an available presentation tool. The extractor cannot read .ppt directly.")
    extractor = EXTRACTORS.get(source.suffix.lower())
    if extractor is None:
        raise ValueError("Supported extraction inputs: PDF, PPTX, DOCX, Markdown, UTF-8 text.")
    if layout_mode not in ("off", "auto") and source.suffix.lower() != ".pdf":
        raise ValueError("--layout applies to PDF sources only.")
    units, review, unit_count = extractor(source)
    layout_pages = parse_layout_pages(layout_mode, units)
    if layout_pages:
        units, review, unit_count = extract_pdf(source, layout_pages)
    merge_units = source.suffix.lower() in MERGED_UNIT_FORMATS
    code_sections, candidates, languages, syntax_checks, recovery_edits = recover_code(
        units, merge_units=merge_units)

    rendered = []
    for unit in units:
        block = unit["heading"] + "\n\n" + unit["body"]
        for candidate in code_sections.get(unit["locator"], []):
            block += "\n\n" + render_candidate_section(candidate)
        rendered.append(block)

    summary = {"document": source.name, "format": source.suffix.lower(), "units": unit_count,
               "review_candidates": review, "images_rendered": 0,
               "layout_pages_extracted": sorted(layout_pages) if layout_pages else [],
               "code_candidates": candidates, "languages": languages,
               "syntax_checks": syntax_checks, "recovery_edits": recovery_edits,
               "targeted_reviews": [],
               "note": ("Review flags are candidates, not a command to render all flagged units. "
                        "Text extraction can miss diagrams, equations, or complex layout even on unflagged pages. "
                        "This extractor renders no images; any later visual review is a separately recorded step. "
                        "Recovered code candidates are recovery assistance, not proof of correct syntax or reading order.")}
    heading = f"# Extracted source: {source.name}\n\nOriginal page, slide, paragraph, and table identifiers below refer to the uploaded source. For text files, cite original section headings or line numbers.\n\n"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(heading + "\n\n".join(rendered) + "\n", encoding="utf-8")
    report_path = output.with_suffix(".extraction.json")
    if report_path.resolve() == source.resolve():
        raise ValueError("The extraction report must not overwrite its source.")
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {**summary, "markdown_path": str(output), "report_path": str(report_path)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_file")
    parser.add_argument("output_markdown")
    parser.add_argument("--layout", default="off",
                        help="PDF page-layout re-extraction: off (default), auto, or page numbers like 4,7")
    arguments = parser.parse_args()
    try:
        print(json.dumps(extract_source(
            arguments.source_file, arguments.output_markdown, arguments.layout),
            ensure_ascii=False))
    except (ValueError, OSError, ImportError) as error:
        parser.exit(1, f"Extraction failed: {error}\n")


if __name__ == "__main__":
    main()
