"""Assemble a schema-version-2 bank into a DOCX (Classic or Card) plus its manifest.

The render manifest is saved beside the DOCX before rendering; a PDF request
converts that finalized DOCX and never reshuffles or redraws content.
"""

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SKILL_DIR = Path(__file__).resolve().parents[1]
if str(SKILL_DIR / "scripts") not in sys.path:
    sys.path.insert(0, str(SKILL_DIR / "scripts"))

import qbank_order  # noqa: E402
import qbank_schema  # noqa: E402
from qbank_formats import card_tokens, formats_root, resolve_format  # noqa: E402
from qbank_render import (add_code_card, add_chip_text, add_options_block,  # noqa: E402
                          convert_docx_to_pdf)

DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "Creating Questions.docx"
SECTION_INSTRUCTIONS = {
    "mcqs": "Choose one answer for each question.",
    "true_false": "Mark each statement True or False.",
    "essays": "Answer every part of each prompt in your own words.",
}
SECTION_TITLES = {"mcqs": "Multiple Choice Questions",
                  "true_false": "True or False Questions",
                  "essays": "Essay Questions"}
LINKED_NOTICE = ("Note: Questions {first} to {last} are linked; later questions "
                 "depend on the results of earlier ones.")


def clear_body(document):
    for child in list(document._body._element):
        if child.tag != qn("w:sectPr"):
            document._body._element.remove(child)


def code_lines(model, block_id):
    return [line["text"] for line in model["code_blocks"][block_id]["lines"]]


def add_stem(document, number, text, tokens, card):
    paragraph = document.add_paragraph(style="List Number")
    num_properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_properties.get_or_add_numId().val = 0
    run = paragraph.add_run(f"{number}. ")
    run.bold = True
    if card:
        add_chip_text(paragraph, text, tokens)
        paragraph.runs[-1].bold = True
    else:
        run = paragraph.add_run(text.strip())
        run.bold = True
    paragraph.paragraph_format.keep_with_next = True


def render_unit(document, model, manifest, unit, section, card, tokens):
    """Render one standalone question or one atomic code group."""
    highlighter = tokens.get("highlighter", "monokai") if card else None
    block_id = unit.get("code_block_id")
    numbers = unit["questions"]
    if block_id:
        block = manifest["blocks"][block_id]
        first = min(numbers.values())
        last = max(numbers.values())
        linked = any(entry["group_id"] == unit["unit_id"]
                     for entry in manifest["linked_groups"])
        if linked:
            document.add_paragraph(LINKED_NOTICE.format(first=first, last=last))
        if card:
            language = model["code_blocks"][block_id]["language"]
            version = block.get("version_hint") or ""
            label = f"Language: {language}{(' ' + version) if version else ''}"
            add_code_card(document, code_lines(model, block_id), language,
                          f"Code Block {block['display_number']} — Questions {first} to {last}",
                          label, tokens, highlighter)
        else:
            for line in code_lines(model, block_id):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(10)
        for question_id in unit["questions"]:
            render_question(document, model, manifest, question_id, card, tokens)
    else:
        render_question(document, model, manifest, next(iter(unit["questions"])),
                        card, tokens)


def render_question(document, model, manifest, question_id, card, tokens):
    question = model["questions"][question_id]
    number = manifest["question_numbers"][question_id]
    stem = question.get("question") or question.get("statement")
    add_stem(document, number, stem, tokens, card)
    if question["section"] == "mcqs":
        if card:
            add_options_block(document, [option["text"] for option in question["options"]],
                              tokens)
        else:
            for letter, option in zip("ABCD", question["options"]):
                paragraph = document.add_paragraph(f"{letter}) {option['text'].strip()}",
                                                   style="List Bullet 2")
                paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
                paragraph.paragraph_format.keep_with_next = letter != "D"
    elif question["section"] == "true_false":
        document.add_paragraph("True / False")
    else:
        document.add_paragraph("." * 80)


def mcq_answer_text(question, manifest):
    letters = manifest["option_letters"][question["id"]]
    letter = letters[question["correct_option_id"]]
    text = next(option["text"].strip() for option in question["options"]
                if option["id"] == question["correct_option_id"])
    return f"{letter}) {text}"


def essay_answer_lines(question):
    lines = [f"• {point.strip()}" for point in question["model_answer"]]
    solution = question.get("reference_solution")
    if solution:
        lines.append("Reference solution:")
        lines.append(solution.strip())
    for case in question.get("test_cases") or []:
        lines.append(f"Case: {case.get('description', '')} → expected: "
                     f"{case.get('expected', '')}")
    return lines


def answer_text(question, manifest):
    if question["section"] == "mcqs":
        return mcq_answer_text(question, manifest)
    if question["section"] == "true_false":
        return question["correct"]
    return "\n".join(essay_answer_lines(question))


def source_label(source, registry, source_mode):
    kind = source["kind"]
    if kind not in {"web", "general"}:
        identifier = registry[(source["document"], kind.removesuffix("_inference"))]
        base_kind = kind.removesuffix("_inference")
        label = f"{identifier}: "
        if base_kind == "pdf":
            label += f"PDF page {source['pdf_page']}"
        elif base_kind == "slides":
            label += f"slide {source['slide']}"
        else:
            if source.get("page"):
                label += f"page {source['page']}, "
            label += source["locator"]
        if source.get("printed_page"):
            label += f", printed page {source['printed_page']}"
        if base_kind != "document" and source.get("locator"):
            label += f", {source['locator']}"
        if kind.endswith("_inference"):
            origin = {"pdf": "PDF", "slides": "slides", "document": "source"}[base_kind]
            label += f"\nInference from {origin}"
        return label
    outside = "Outside PDF" if source_mode.startswith("pdf_") else "Outside uploaded material"
    if kind == "web":
        return f"{outside}: Web: {source['title']}\n{source['url']}"
    return (f"{outside}: General knowledge from model; checked against "
            f"{source['verified_against']}")


def fill_answer_cells(cells, content):
    for cell, text in zip(cells, content):
        cell.text = text
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(10)


def add_answer_key(document, model, manifest, sections_order):
    document.add_page_break()
    document.add_paragraph("Model Answer", style="Heading 1")
    registry = {identifier: record for identifier, record in manifest["sources"].items()}
    legend_registry = {(record["document"], record["kind"]): identifier
                       for identifier, record in registry.items()}
    if legend_registry:
        document.add_paragraph(
            "Source references below identify the uploaded files. PDF pages and slides "
            "use 1-based numbering. Inference labels identify applications or deductions "
            "from the cited material.")
        for (document_name, _), identifier in sorted(legend_registry.items(),
                                                     key=lambda item: item[1]):
            document.add_paragraph(f"{identifier}: {document_name}")
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    table.autofit = False
    for column, width in zip(table.columns, (0.45, 1.65, 1.8, 2.1)):
        column.width = Inches(width)
    header = table.rows[0]
    for cell, column in zip(header.cells, table.columns):
        cell.width = column.width
    header._tr.get_or_add_trPr().append(qbank_schema_header_element())
    fill_answer_cells(header.cells, ["No.", "Answer", "Explanation or correction", "Source"])
    for cell in header.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for section in qbank_schema.QUESTION_TYPES:
        for question_id in sections_order[section]:
            question = model["questions"][question_id]
            number = manifest["question_numbers"][question_id]
            explanation = question["explanation"].strip()
            if section == "true_false" and question["correct"] == "False":
                explanation = "Correction: " + question["correction"].strip()
            labels = list(dict.fromkeys(
                source_label(source, legend_registry, model["source_mode"])
                for source in question["sources"]))
            fill_answer_cells(table.add_row().cells,
                              [str(number), answer_text(question, manifest),
                               explanation, "\n".join(labels)])
    issues = model.get("source_issues") or []
    if issues:
        document.add_paragraph("Source issues", style="Heading 2")
        for issue in issues:
            document.add_paragraph(str(issue), style="List Bullet")


def qbank_schema_header_element():
    from docx.oxml import OxmlElement
    return OxmlElement("w:tblHeader")


def build_v2_docx(bank, output_path, template_path=DEFAULT_TEMPLATE,
                  format_choice=None, variant=None):
    """Validate, finalize the order once, save the manifest, and render."""
    model = qbank_schema.normalize_bank(bank)
    sections_order, manifest = qbank_order.finalize_order(model, seed=model.get("shuffle_seed"))
    chosen = format_choice or (model.get("format") or {}).get("profile_id")
    if chosen in (None, "", "classic"):
        chosen = "classic"
        resolved_variant = None
    else:
        resolved = resolve_format(chosen, variant or (model.get("format") or {}).get("variant"),
                                  formats_root(SKILL_DIR))
        chosen = resolved["profile"]["id"]
        resolved_variant = resolved["variant"]
    card = chosen == "builtin-card"
    tokens = card_tokens(resolve_format("builtin-card", resolved_variant)["profile"],
                         resolved_variant) if card else None

    output_path = Path(output_path)
    manifest_path = output_path.with_suffix(".manifest.json")
    manifest["format"] = {"profile_id": chosen, "variant": resolved_variant}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")

    document = Document(str(template_path))
    clear_body(document)
    title = document.add_paragraph(model["title"].strip(), style="Title")
    document.add_paragraph(
        f"{sum(model['counts'].values())} questions: {model['counts']['mcqs']} multiple "
        f"choice, {model['counts']['true_false']} True/False, and "
        f"{model['counts']['essays']} essay.")
    for section in qbank_schema.QUESTION_TYPES:
        units = manifest["sections"][section]
        if not units:
            continue
        heading = document.add_paragraph(SECTION_TITLES[section], style="Heading 1")
        if section != "mcqs":
            heading.paragraph_format.page_break_before = True
        instruction = document.add_paragraph(SECTION_INSTRUCTIONS[section])
        instruction.paragraph_format.keep_with_next = True
        for unit in units:
            render_unit(document, model, manifest, unit, section, card, tokens)
    if model.get("include_answer_key", True):
        add_answer_key(document, model, manifest, sections_order)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return {"manifest_path": str(manifest_path), "output_path": str(output_path),
            "format": manifest["format"]}


def build_pdf_from_docx(docx_path, pdf_path):
    """Convert the finalized DOCX; never redraw or reshuffle."""
    return convert_docx_to_pdf(docx_path, pdf_path)
