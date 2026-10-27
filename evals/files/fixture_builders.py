"""Builders for synthetic public test fixtures: text PDFs, code slides, code documents."""

import textwrap


def escape_pdf_text(text):
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_text_pdf(path, pages):
    """Write a minimal valid PDF whose pages show the given text lines.

    `pages` is a list of line lists. Uses the built-in Courier font so no
    external dependency or binary fixture is needed.
    """
    objects = {}
    page_ids = []
    font_id = 3 + 2 * len(pages)
    for index, lines in enumerate(pages):
        page_id = 3 + 2 * index
        content_id = page_id + 1
        page_ids.append(page_id)
        content = ["BT", "/F1 10 Tf", "72 720 Td", "12 TL"]
        for line in lines:
            content.append(f"({escape_pdf_text(line)}) Tj T*")
        content.append("ET")
        stream = "\n".join(content).encode("latin-1", "replace")
        objects[page_id] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                            f"/Contents {content_id} 0 R >>")
        objects[content_id] = (f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1")
                               + stream + b"\nendstream")
    objects[1] = "<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>"
    objects[font_id] = "<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        body = objects[number]
        if isinstance(body, str):
            body = body.encode("latin-1", "replace")
        out += f"{number} 0 obj\n".encode("latin-1") + body + b"\nendobj\n"
    xref_position = len(out)
    out += f"xref\n0 {font_id + 1}\n".encode("latin-1")
    out += b"0000000000 65535 f \n"
    for number in sorted(objects):
        out += f"{offsets[number]:010d} 00000 n \n".encode("latin-1")
    out += (f"trailer\n<< /Size {font_id + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_position}\n%%EOF\n").encode("latin-1")
    path.write_bytes(bytes(out))


def add_code_textbox(slide, lines, left_inches=1.0, top_inches=1.0):
    """Add a textbox holding one paragraph per given code line."""
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(left_inches), Inches(top_inches),
                                   Inches(6), Inches(0.3 * max(len(lines), 1)))
    frame = box.text_frame
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
    return box


COUNTER_CODE = textwrap.dedent("""\
    class Counter:
        def __init__(self, value):
            self.value = value
        def advance(self):
            self.value += 3
            return self.value
    c = Counter(4)
    print(c.value)
    """).splitlines()
