"""Build the Gate A copy/conversion prototype DOCX (synthetic, public)."""

import sys
from pathlib import Path

SKILL_SCRIPTS = Path(__file__).resolve().parents[3] / "scripts"
if str(SKILL_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SKILL_SCRIPTS))

from docx import Document  # noqa: E402

from qbank_formats import card_tokens, resolve_format  # noqa: E402
from qbank_render import add_code_card, add_options_block  # noqa: E402

NINE_LINE_EXAMPLE = """\
class Counter:
    def __init__(self, value):
        self.value = value

    def advance(self):
        self.value += 3
        return self.value

c = Counter(4)""".splitlines()

TRICKY_EXAMPLE = """\
def report(scores, label):
    total = 0
    for index, score in enumerate(scores, 1):
        adjusted = score + index_1_bonus
        if __debug__:
            print(f"{label}: student_{index} -> {adjusted}")
        total += adjusted
    message = "quotes 'inside' and \\"double\\" plus unicode: α-β-γ"
    long_line = message + " padded to exceed the printed width so the renderer must handle a very long single logical line without corrupting it"
    return total, message, long_line""".splitlines()


def hundred_line_example():
    lines = ["# 100-line gutter calibration listing", "values = []"]
    for index in range(1, 98):
        lines.append(f"values.append({index} * item_scale_{index % 5 + 1})")
    lines.append("print(sum(values))")
    return lines


def main():
    output = Path(__file__).resolve().parent / "copy_prototype.docx"
    resolved = resolve_format("builtin-card", "monokai")
    tokens = card_tokens(resolved["profile"], "monokai")
    document = Document(str(SKILL_SCRIPTS.parent / "assets" / "Creating Questions.docx"))
    body = document.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)
    document.add_paragraph("Copy and conversion prototype", style="Title")
    document.add_paragraph(
        "Synthetic public fixture for Gate A: select the code column (click in the first "
        "code line, extend to the last), copy, and paste into an editor. Numbers live in "
        "a separate gutter column and are not part of the code text.")

    add_code_card(document, NINE_LINE_EXAMPLE, "python",
                  "Code Block 1 — Nine lines", "Language: Python 3.x", tokens, "monokai")
    document.add_paragraph()
    add_code_card(document, TRICKY_EXAMPLE, "python",
                  "Code Block 2 — Underscores, quotes, Unicode, long line",
                  "Language: Python 3.x", tokens, "monokai")
    document.add_paragraph()
    add_code_card(document, hundred_line_example(), "python",
                  "Code Block 3 — One hundred lines", "Language: Python 3.x",
                  tokens, "monokai")

    document.add_paragraph("1) Which value does c.value hold at the end of Code Block 1?")
    add_options_block(document, ["`4`", "`7`", "`10`", "`AttributeError`"], tokens)
    document.add_paragraph("2) Which statement matches the long option layout rule?")
    add_options_block(
        document,
        ["`total` accumulates the adjusted scores and is returned with the message",
         "The function prints the final total and returns nothing",
         "The long line is split into two executable statements by the renderer",
         "Unicode characters in the string are replaced with placeholders"],
        tokens)

    document.save(str(output))
    print(output)


if __name__ == "__main__":
    main()
