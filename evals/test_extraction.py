import importlib.util
import json
import shutil
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter


SCRIPT = Path(__file__).resolve().parents[1] / 'scripts/extract_source_markdown.py'
SPEC = importlib.util.spec_from_file_location('extractor', SCRIPT)
extractor = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(extractor)


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.scratch = Path.cwd() / 'work' / ('extraction_' + self.id().split('.')[-1])
        self.scratch.mkdir(parents=True, exist_ok=True)
        self.output = self.scratch / 'extracted.md'
        self.addCleanup(shutil.rmtree, self.scratch)

    def test_pdf_retains_page_indices_and_flags_missing_text_without_images(self):
        source = self.scratch / 'scan.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        writer.write(source)
        before = source.read_bytes()
        report = extractor.extract_source(source, self.output)
        markdown = self.output.read_text(encoding='utf-8')
        self.assertIn('## PDF page 1', markdown)
        self.assertIn('## PDF page 2', markdown)
        self.assertEqual([flag['pdf_page'] for flag in report['review_candidates']], [1, 2])
        self.assertEqual(report['images_rendered'], 0)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(json.loads(self.output.with_suffix('.extraction.json').read_text())['units'], 2)
        self.assertFalse(list(self.scratch.glob('*.png')))

    def test_pptx_preserves_slide_text_tables_and_notes(self):
        source = self.scratch / 'lecture.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Study population'
        slide.notes_slide.notes_text_frame.text = 'Only the sampled population was studied.'
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for row, values in zip(table.rows, [('Group', 'Count'), ('A', '12')]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(source)
        report = extractor.extract_source(source, self.output)
        markdown = self.output.read_text(encoding='utf-8')
        self.assertIn('## Slide 1', markdown)
        self.assertIn('Study population', markdown)
        self.assertIn('| A | 12 |', markdown)
        self.assertIn('Only the sampled population was studied.', markdown)
        self.assertEqual(report['units'], 2)
        self.assertEqual(report['review_candidates'][0]['slide'], 2)
        self.assertEqual(report['images_rendered'], 0)

    def test_docx_preserves_body_order_with_document_locators(self):
        source = self.scratch / 'paper.docx'
        paper = Document()
        paper.add_heading('Methods', level=1)
        table = paper.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Variable'
        table.cell(0, 1).text = 'Value'
        table.cell(1, 0).text = 'Sample'
        table.cell(1, 1).text = '30'
        paper.add_paragraph('The study did not establish causation.')
        paper.save(source)
        extractor.extract_source(source, self.output)
        markdown = self.output.read_text(encoding='utf-8')
        self.assertLess(markdown.index('Methods'), markdown.index('| Sample | 30 |'))
        self.assertLess(markdown.index('| Sample | 30 |'), markdown.index('did not establish causation'))
        self.assertIn('## Paragraph 1', markdown)
        self.assertIn('## Paragraph 2', markdown)
        self.assertIn('## Table 1', markdown)
        self.assertNotIn('PDF page', markdown)

    def test_markdown_retains_original_headings_and_unicode(self):
        source = self.scratch / 'paper.md'
        original = '# Results\n\nThe mean is 12; this is not a causal result.\nGreek symbol: α.\n'
        source.write_text(original, encoding='utf-8')
        report = extractor.extract_source(source, self.output)
        self.assertIn(original, self.output.read_text(encoding='utf-8'))
        self.assertEqual(report['review_candidates'], [])
        self.assertEqual(report['images_rendered'], 0)

    def test_source_cannot_be_overwritten_by_extraction(self):
        source = self.scratch / 'notes.md'
        source.write_text('Keep this source.', encoding='utf-8')
        with self.assertRaises(ValueError):
            extractor.extract_source(source, source)
        self.assertEqual(source.read_text(), 'Keep this source.')

    def test_legacy_ppt_requires_conversion_and_creates_no_false_output(self):
        source = self.scratch / 'old.ppt'
        source.write_bytes(b'legacy-placeholder')
        with self.assertRaisesRegex(ValueError, 'conversion'):
            extractor.extract_source(source, self.output)
        self.assertFalse(self.output.exists())
        self.assertEqual(source.read_bytes(), b'legacy-placeholder')


if __name__ == '__main__':
    unittest.main()
