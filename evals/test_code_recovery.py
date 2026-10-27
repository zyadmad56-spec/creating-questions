"""Phase 2 acceptance checks: code recovery, language detection, and honest review states."""

import importlib.util
import shutil
import sys
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation

EVALS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(EVALS_DIR / "files"))
from fixture_builders import COUNTER_CODE, add_code_textbox, make_text_pdf  # noqa: E402

SCRIPT = EVALS_DIR.parent / "scripts" / "extract_source_markdown.py"
SPEC = importlib.util.spec_from_file_location("extractor", SCRIPT)
extractor = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(extractor)


class CodeRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.scratch = Path.cwd() / "work" / ("recovery_" + self.id().split(".")[-1])
        self.scratch.mkdir(parents=True, exist_ok=True)
        self.output = self.scratch / "extracted.md"
        self.addCleanup(shutil.rmtree, self.scratch)

    def extract(self, source, layout="off"):
        report = extractor.extract_source(source, self.output, layout)
        markdown = self.output.read_text(encoding="utf-8")
        return report, markdown

    def test_pdf_code_page_is_recovered_with_language_and_locator(self):
        source = self.scratch / "lecture.pdf"
        make_text_pdf(source, [
            ["Introduction to loops.", "A loop repeats a block of code."],
            ["Example program:", ""] + COUNTER_CODE,
        ])
        report, markdown = self.extract(source)
        candidates = report["code_candidates"]
        self.assertTrue(candidates, "The python listing must be detected")
        counter = next(c for c in candidates if "class Counter" in c["normalized_code"])
        self.assertEqual(counter["language"], "python")
        self.assertEqual(counter["locator"], "PDF page 2")
        self.assertIn(counter["review_status"], {"recovered", "needs_review"})
        self.assertGreaterEqual(counter["language_confidence"], 2)
        self.assertIn("```python", markdown)
        self.assertIn("Recovered code candidate", markdown)
        self.assertIn("## PDF page 2", markdown)
        self.assertEqual(report["images_rendered"], 0)

    def test_layout_mode_recovers_indented_code_on_requested_pages_only(self):
        source = self.scratch / "lecture.pdf"
        make_text_pdf(source, [["Plain prose page with words."]] + [COUNTER_CODE])
        report, _ = self.extract(source, layout="2")
        self.assertEqual(report["layout_pages_extracted"], [2])
        layout_candidates = [c for c in report["code_candidates"]
                             if c["locator"] == "PDF page 2 (layout extraction)"]
        self.assertTrue(layout_candidates, "Layout text must be scanned with its own locator")
        recovered = next(c for c in layout_candidates if c["language"] == "python"
                         and c["review_status"] == "recovered")
        self.assertEqual(recovered["normalized_code"].splitlines()[0], "class Counter:")
        self.assertIn("    def __init__(self, value):", recovered["normalized_code"])
        self.assertTrue(all(c["locator"] != "PDF page 1 (layout extraction)"
                            for c in report["code_candidates"]))

    def test_numbered_display_labels_are_mapped_and_removed(self):
        source = self.scratch / "exam.txt"
        numbered = [f"{number:02d}. {line}"
                    for number, line in enumerate(COUNTER_CODE[:5], 1)]
        source.write_text("Model answer listing:\n" + "\n".join(numbered) + "\n", encoding="utf-8")
        report, _ = self.extract(source)
        candidate = report["code_candidates"][0]
        self.assertEqual(candidate["normalized_code"].splitlines()[0], "class Counter:")
        self.assertEqual(candidate["line_labels"][0], {"display_label": 1, "line_index": 0})
        self.assertEqual(candidate["line_labels"][-1]["display_label"], 5)
        self.assertTrue(any(edit["edit"] == "removed printed display-number prefixes"
                            for edit in candidate["recovery_edits"]))

    def test_pptx_code_shape_keeps_paragraph_lines_and_blank_lines(self):
        source = self.scratch / "lecture.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        code_lines = COUNTER_CODE[:4] + [""] + COUNTER_CODE[4:6]
        add_code_textbox(slide, code_lines)
        deck.save(source)
        report, _ = self.extract(source)
        candidate = next(c for c in report["code_candidates"] if c["language"] == "python")
        self.assertEqual(candidate["normalized_code"].splitlines(), code_lines)
        self.assertEqual(candidate["locator"], "Slide 1")
        self.assertNotIn("\r", candidate["normalized_code"])

    def test_docx_consecutive_code_paragraphs_merge_into_one_candidate(self):
        source = self.scratch / "handout.docx"
        paper = Document()
        paper.add_paragraph("The class below stores two values.")
        for line in COUNTER_CODE[:5]:
            paper.add_paragraph(line)
        paper.add_paragraph("Run it to see the stored value.")
        paper.save(source)
        report, _ = self.extract(source)
        merged = [c for c in report["code_candidates"] if c["source_kind"] == "consecutive_paragraphs"]
        self.assertTrue(merged, "Consecutive code paragraphs must merge into one candidate")
        self.assertEqual(len(merged[0]["normalized_code"].splitlines()), 5)
        self.assertEqual(merged[0]["locator"], "Paragraph 2")

    def test_intentional_syntax_error_is_recorded_without_rewriting(self):
        source = self.scratch / "broken.txt"
        source.write_text("def broken(:\n    return 1\n    return 2\n    return 3\n",
                          encoding="utf-8")
        report, _ = self.extract(source)
        candidate = next(c for c in report["code_candidates"] if c["language"] == "python")
        check = next(c for c in report["syntax_checks"] if c["result"] == "syntax_error")
        self.assertEqual(candidate["candidate_id"], check["candidate_id"])
        self.assertEqual(candidate["review_status"], "needs_review")
        self.assertEqual(candidate["original_text"], "def broken(:\n    return 1\n    return 2\n    return 3")
        self.assertIn("May be intentional", check["note"])

    def test_fill_the_gap_snippet_is_preserved_and_flagged(self):
        source = self.scratch / "gap.txt"
        source.write_text("class CourseRecord:\n    def __init__(self, hours):\n"
                          "        self.hours = hours\n\n    def read_hours(self):\n"
                          "        return ______\n    def extra(self):\n        return self.hours\n",
                          encoding="utf-8")
        report, markdown = self.extract(source)
        candidate = next(c for c in report["code_candidates"] if "______" in c["normalized_code"])
        self.assertEqual(candidate["review_status"], "needs_review")
        self.assertIn("______", markdown)

    def test_multiple_languages_are_detected_with_constructs(self):
        source = self.scratch / "mixed.txt"
        source.write_text(
            "query:\nSELECT name FROM students WHERE grade > 3;\nSELECT name FROM courses;\n"
            "SELECT count FROM totals;\n\n"
            "utility:\n#include <stdio.h>\nint main(void) {\n    printf(\"hi\");\n    return 0;\n}\n"
            "int other(void) {\n    return 1;\n}\n",
            encoding="utf-8")
        report, _ = self.extract(source)
        self.assertIn("sql", report["languages"])
        self.assertIn("c", report["languages"])
        self.assertIn("select_queries", report["languages"]["sql"]["constructs"])
        self.assertIn("printf_output", report["languages"]["c"]["constructs"])

    def test_no_code_document_yields_no_candidates(self):
        source = self.scratch / "prose.docx"
        paper = Document()
        paper.add_paragraph("A triangle has three sides. A square has four equal sides.")
        paper.add_paragraph("The study did not establish causation.")
        paper.save(source)
        report, _ = self.extract(source)
        self.assertEqual(report["code_candidates"], [])
        self.assertEqual(report["languages"], {})
        self.assertEqual(report["syntax_checks"], [])

    def test_dense_flattened_code_page_is_flagged_beyond_the_word_count_heuristic(self):
        source = self.scratch / "dense.pdf"
        make_text_pdf(source, [[
            "def summarize(values): total = 0 for value in values: total += value",
            "return total / len(values) print(summarize([1, 2, 3])) print(summarize([4, 5]))",
            "print(summarize([7])) print(summarize([]))",
        ]])
        report, _ = self.extract(source)
        self.assertTrue(report["code_candidates"],
                        "A dense code page must be flagged even with plenty of words")

    def test_source_bytes_and_locators_survive_extraction(self):
        source = self.scratch / "lecture.pdf"
        make_text_pdf(source, [["Words on page one."] * 4, COUNTER_CODE])
        before = source.read_bytes()
        report, markdown = self.extract(source)
        self.assertEqual(source.read_bytes(), before)
        self.assertIn("## PDF page 1", markdown)
        self.assertIn("## PDF page 2", markdown)
        self.assertEqual(report["units"], 2)
        self.assertEqual(report["images_rendered"], 0)

    def test_markdown_fenced_block_language_label_is_respected(self):
        source = self.scratch / "notes.md"
        source.write_text("Demo:\n\n```sql\nSELECT id FROM users;\nSELECT name FROM users;\n```\n",
                          encoding="utf-8")
        report, markdown = self.extract(source)
        candidate = report["code_candidates"][0]
        self.assertEqual(candidate["language"], "sql")
        self.assertEqual(candidate["stated_language"], "sql")
        self.assertIn("```sql", markdown)

    def test_layout_flag_is_rejected_for_non_pdf_sources(self):
        source = self.scratch / "notes.txt"
        source.write_text("plain text\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "PDF sources only"):
            self.extract(source, layout="1,2")


if __name__ == "__main__":
    unittest.main()
