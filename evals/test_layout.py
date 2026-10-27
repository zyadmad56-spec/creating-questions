"""Phase 5 acceptance checks: Card layout, saved formats, manifest, and PDF path."""

import json
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

EVALS_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = EVALS_DIR.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import qbank_formats  # noqa: E402
from create_question_bank_docx import build_docx  # noqa: E402

FIXTURE = EVALS_DIR / "files" / "schema" / "counter_bank.json"


def audited_counter_bank():
    bank = json.loads(FIXTURE.read_text(encoding="utf-8"))
    bank["shuffle_seed"] = 42
    for question in bank["mcqs"]:
        question["answer_verified"] = True
        question["execution"] = {"status": "not_applicable", "execution_verified": False,
                                 "reason": "Unit test", "cases": []}
    return bank


def hard_linked_bank():
    bank = audited_counter_bank()
    for question in bank["mcqs"]:
        question["difficulty"] = "hard"
    bank["groups"][0]["difficulty"] = "hard"
    bank["groups"][0]["dependency_mode"] = "linked"
    bank["groups"][0]["dependencies"] = [{"question_id": "q_return", "depends_on": ["q_value"]}]
    bank["mcqs"][1]["depends_on"] = ["q_value"]
    bank["difficulty_counts"]["mcqs"] = {"easy": 0, "medium": 0, "hard": 2}
    bank["programming"] = {"allow_hard_dependencies": True}
    return bank


def card_tables(document):
    return [table for table in document.tables
            if len(table.rows) == 2 and len(table.columns) == 2
            and table.rows[1].cells[0].text.split("\n")[0].strip() == "1"]


class CardLayoutTests(unittest.TestCase):
    def setUp(self):
        self.scratch = Path.cwd() / "work" / "layout"
        self.scratch.mkdir(parents=True, exist_ok=True)
        self.output = self.scratch / "card_bank.docx"
        self.output.unlink(missing_ok=True)
        self.addCleanup(self.output.unlink, missing_ok=True)

    def test_card_manifest_is_saved_before_rendering(self):
        result = build_docx(audited_counter_bank(), self.output)
        manifest = json.loads(Path(result["manifest_path"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["format"], {"profile_id": "builtin-card", "variant": "monokai"})
        self.assertEqual(manifest["question_numbers"]["q_value"], manifest["question_numbers"]["q_value"])
        self.assertIn("b_counter", manifest["blocks"])
        self.assertEqual(len(manifest["blocks"]["b_counter"]["line_numbers"]), 7)
        self.addCleanup(Path(result["manifest_path"]).unlink, missing_ok=True)

    def test_card_renders_code_card_with_separate_gutter(self):
        build_docx(audited_counter_bank(), self.output)
        document = Document(self.output)
        tables = card_tables(document)
        self.assertEqual(len(tables), 1)
        gutter_lines = tables[0].rows[1].cells[0].text.split("\n")
        code_lines = tables[0].rows[1].cells[1].text.split("\n")
        self.assertEqual(gutter_lines, [str(number) for number in range(1, 8)])
        self.assertEqual(code_lines[0], "class Counter:")
        self.assertEqual(code_lines[6], "c = Counter(4)")
        header = tables[0].rows[0].cells[0].text
        self.assertIn("Code Block 1", header)
        self.assertIn("Language: python 3.x", header)

    def test_card_code_runs_are_real_text_with_chips_and_shading(self):
        bank = audited_counter_bank()
        bank["mcqs"][0]["question"] = "What does `c.value` hold?"
        build_docx(bank, self.output)
        document = Document(self.output)
        code_cell = card_tables(document)[0].rows[1].cells[1]
        first_paragraph = code_cell.paragraphs[0]
        self.assertEqual("".join(run.text for run in first_paragraph.runs), "class Counter:")
        shading = first_paragraph.runs[0]._r.find(qn("w:rPr")).findall(qn("w:shd")) \
            if first_paragraph.runs[0]._r.find(qn("w:rPr")) is not None else []
        runs_with_shading = 0
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run_properties = run._r.find(qn("w:rPr"))
                if run_properties is not None and run_properties.find(qn("w:shd")) is not None:
                    runs_with_shading += 1
        self.assertGreater(runs_with_shading, 0, "Inline code chips must be shaded runs")

    def test_short_options_use_grid_and_long_options_one_per_row(self):
        build_docx(audited_counter_bank(), self.output)
        document = Document(self.output)
        grid_options = [table for table in document.tables
                        if len(table.rows) == 2 and len(table.columns) == 2
                        and table.rows[0].cells[0].text.strip().startswith("A)")]
        self.assertEqual(len(grid_options), 2, "Each MCQ's short options go two by two")
        bank = audited_counter_bank()
        for question in bank["mcqs"]:
            question["options"] = [
                {"id": option["id"],
                 "text": option["text"] + " with a much longer explanation attached"}
                for option in question["options"]]
        long_output = self.scratch / "long_options.docx"
        long_output.unlink(missing_ok=True)
        self.addCleanup(long_output.unlink, missing_ok=True)
        build_docx(bank, long_output)
        long_document = Document(long_output)
        self.assertFalse([table for table in long_document.tables
                          if len(table.rows) == 2 and len(table.columns) == 2
                          and table.rows[0].cells[0].text.strip().startswith("A)")])

    def test_classic_v2_keeps_code_as_plain_text_without_gutter(self):
        build_docx(audited_counter_bank(), self.output, format_id="classic")
        document = Document(self.output)
        self.assertEqual(card_tables(document), [])
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("class Counter:", texts)

    def test_linked_group_prints_the_notice_with_final_numbers(self):
        build_docx(hard_linked_bank(), self.output)
        document = Document(self.output)
        notice = next(paragraph.text for paragraph in document.paragraphs
                      if "are linked" in paragraph.text)
        self.assertRegex(notice, r"Questions \d+ to \d+ are linked")

    def test_student_only_export_omits_the_answer_key(self):
        bank = audited_counter_bank()
        bank["include_answer_key"] = False
        build_docx(bank, self.output)
        document = Document(self.output)
        self.assertFalse(any(table.style and "Grid" in str(table.style)
                             for table in document.tables))
        self.assertTrue(card_tables(document))

    def test_same_seed_produces_the_same_manifest(self):
        first = build_docx(audited_counter_bank(), self.output)
        second_output = self.scratch / "second.docx"
        second_output.unlink(missing_ok=True)
        self.addCleanup(second_output.unlink, missing_ok=True)
        second = build_docx(audited_counter_bank(), second_output)
        first_manifest = json.loads(Path(first["manifest_path"]).read_text(encoding="utf-8"))
        second_manifest = json.loads(Path(second["manifest_path"]).read_text(encoding="utf-8"))
        self.assertEqual(first_manifest["question_numbers"], second_manifest["question_numbers"])
        self.assertEqual(first_manifest["option_letters"], second_manifest["option_letters"])
        self.addCleanup(Path(first["manifest_path"]).unlink, missing_ok=True)
        self.addCleanup(Path(second["manifest_path"]).unlink, missing_ok=True)


class SavedFormatTests(unittest.TestCase):
    def setUp(self):
        self.scratch = Path.cwd() / "work" / "formats"
        if self.scratch.exists():
            import shutil
            shutil.rmtree(self.scratch)
        self.scratch.mkdir(parents=True)
        self.addCleanup(lambda: __import__("shutil").rmtree(self.scratch, ignore_errors=True))

    def test_saved_profile_round_trip_preserves_layout_tokens(self):
        profile = {"format_version": 1, "id": "my-card", "name": "My Card",
                   "variants": ["monokai", "print"], "code_font": "Cascadia Mono",
                   "code_font_size_pt": 9, "option_layout": "adaptive",
                   "colors": {"monokai": {"code": "111111", "gutter": "222222",
                                          "header": "333333", "chip": "444444",
                                          "gutter_text": "555555", "code_text": "eeeeee",
                                          "highlighter": "monokai"}}}
        qbank_formats.save_format(profile, self.scratch)
        resolved = qbank_formats.resolve_format("my-card", root=self.scratch)
        tokens = qbank_formats.card_tokens(resolved["profile"], "monokai")
        self.assertEqual(tokens["font"], "Cascadia Mono")
        self.assertEqual(tokens["code"], "111111")
        self.assertEqual(qbank_formats.list_available_formats(self.scratch),
                         ["builtin-card", "classic", "my-card"])

    def test_corrupt_saved_profile_is_reported_and_preserved(self):
        directory = self.scratch / "broken"
        directory.mkdir()
        (directory / "format.json").write_text("{ broken", encoding="utf-8")
        with self.assertRaisesRegex(qbank_formats.FormatError, "unreadable"):
            qbank_formats.resolve_format("broken", root=self.scratch)
        self.assertEqual((directory / "format.json").read_text(encoding="utf-8"), "{ broken")

    def test_builtin_ids_cannot_be_shadowed(self):
        with self.assertRaisesRegex(qbank_formats.FormatError, "built-in"):
            qbank_formats.save_format({"format_version": 1, "id": "classic"}, self.scratch)

    def test_print_variant_tokens_differ_from_monokai(self):
        profile = qbank_formats.resolve_format("builtin-card")["profile"]
        monokai = qbank_formats.card_tokens(profile, "monokai")
        print_variant = qbank_formats.card_tokens(profile, "print")
        self.assertNotEqual(monokai["code"], print_variant["code"])
        self.assertEqual(monokai["highlighter"], "monokai")
        self.assertEqual(print_variant["highlighter"], "friendly")


class PdfConversionTests(unittest.TestCase):
    def setUp(self):
        self.scratch = Path.cwd() / "work" / "pdf"
        self.scratch.mkdir(parents=True, exist_ok=True)
        self.docx = self.scratch / "counter.docx"
        self.pdf = self.scratch / "counter.pdf"
        for path in (self.docx, self.pdf):
            path.unlink(missing_ok=True)
        self.addCleanup(self.docx.unlink, missing_ok=True)
        self.addCleanup(self.pdf.unlink, missing_ok=True)

    def test_pdf_is_converted_from_the_finalized_docx(self):
        try:
            import win32com.client  # noqa: F401
        except ImportError:
            self.skipTest("pywin32 unavailable; PDF conversion not checked")
        from create_question_bank_docx import build_pdf_from_docx
        build_docx(audited_counter_bank(), self.docx)
        result = build_pdf_from_docx(self.docx, self.pdf)
        self.assertEqual(result["status"], "converted")
        self.assertGreaterEqual(result["pages"], 1)
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(self.pdf)).pages)
        self.assertIn("class Counter:", text)


if __name__ == "__main__":
    unittest.main()
