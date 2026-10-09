import copy
import importlib.util
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPT_PATH = SKILL_DIR / "scripts" / "create_question_bank_docx.py"
SPEC = importlib.util.spec_from_file_location("formatter", SCRIPT_PATH)
formatter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(formatter)


def sample_bank():
    pdf_source = {
        "kind": "pdf", "document": "practice.pdf", "pdf_page": 1,
        "evidence": "A triangle has three sides. A square has four equal sides."
    }
    shared = {"explanation": "Compare the number of sides.",
              "sources": [pdf_source], "answer_verified": True}
    return {
        "title": "Geometry Practice Questions", "source_mode": "pdf_only",
        "counts": {"mcqs": 2, "true_false": 2, "essays": 2},
        "difficulty_counts": {
            "mcqs": {"easy": 1, "medium": 0, "hard": 1},
            "true_false": {"easy": 1, "medium": 1, "hard": 0},
            "essays": {"easy": 0, "medium": 1, "hard": 1}
        },
        "mcqs": [
            {**copy.deepcopy(shared), "question": "How many sides does a triangle have?",
             "options": ["Two sides", "Three sides", "Four sides", "Five sides"],
             "correct": "B", "difficulty": "easy"},
            {**copy.deepcopy(shared), "question": "Which description fits a square?",
             "options": ["Four equal sides", "Three equal sides", "Two open sides", "Five open sides"],
             "correct": "A", "difficulty": "hard"}
        ],
        "true_false": [
            {**copy.deepcopy(shared), "statement": "A triangle does not have four sides.",
             "correct": "True", "difficulty": "easy"},
            {**copy.deepcopy(shared), "statement": "A square has three sides.",
             "correct": "False", "correction": "A square has four equal sides.",
             "difficulty": "medium"}
        ],
        "essays": [
            {**copy.deepcopy(shared), "question": "Compare triangle and square side counts.",
             "model_answer": ["A triangle has three sides.", "A square has four equal sides."],
             "difficulty": "medium"},
            {**copy.deepcopy(shared), "question": "Explain why a four-sided shape is not a triangle.",
             "model_answer": ["A triangle must have three sides."], "difficulty": "hard"}
        ]
    }


class FormatterTests(unittest.TestCase):
    def setUp(self):
        self.bank = sample_bank()
        scratch = Path.cwd() / "work"
        scratch.mkdir(exist_ok=True)
        self.output = scratch / (self.id().split(".")[-1] + ".docx")
        self.output.unlink(missing_ok=True)
        self.addCleanup(self.output.unlink, missing_ok=True)

    def test_default_output_includes_essays_and_matching_global_numbers(self):
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        stems = [p.text for p in document.paragraphs if p.style.name == "List Number"]
        self.assertEqual(len(stems), 6)
        self.assertEqual([s.split(".", 1)[0] for s in stems], ["1", "2", "3", "4", "5", "6"])
        self.assertTrue(any(p.text == "Model Answer" for p in document.paragraphs))
        self.assertEqual([r.cells[0].text for r in document.tables[0].rows[1:]],
                         ["1", "2", "3", "4", "5", "6"])
        self.assertTrue(any("A square has four equal sides." in row.cells[1].text
                            for row in document.tables[0].rows[5:]))

    def test_correct_answer_tracks_option_shuffling(self):
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        paragraphs = [p.text for p in document.paragraphs]
        self.assertTrue(document.tables, "The default answer key must be present")
        for row in document.tables[0].rows[1:3]:
            number = row.cells[0].text
            stem_position = next(i for i, text in enumerate(paragraphs) if text.startswith(number + ". "))
            answer_text = row.cells[1].text
            letter, correct_text = answer_text.split(") ", 1)
            self.assertEqual(paragraphs[stem_position + "ABCD".index(letter) + 1], answer_text)
            expected = "Three sides" if "triangle" in paragraphs[stem_position] else "Four equal sides"
            self.assertEqual(correct_text, expected)

    def test_template_page_geometry_is_preserved(self):
        formatter.build_docx(self.bank, self.output)
        original = Document(SKILL_DIR / "assets" / "Creating Questions.docx").sections[0]
        generated = Document(self.output).sections[0]
        self.assertEqual((generated.top_margin, generated.bottom_margin, generated.left_margin,
                          generated.right_margin, generated.page_width, generated.page_height),
                         (original.top_margin, original.bottom_margin, original.left_margin,
                         original.right_margin, original.page_width, original.page_height))

    def test_answer_table_header_uses_same_widths_as_its_body(self):
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        table = document.tables[0]
        self.assertEqual([cell.width for cell in table.rows[0].cells],
                         [cell.width for cell in table.rows[1].cells])
        section = document.sections[0]
        self.assertLessEqual(sum(cell.width for cell in table.rows[1].cells),
                             section.page_width - section.left_margin - section.right_margin)

    def test_bad_counts_are_rejected_before_writing(self):
        self.bank["counts"]["mcqs"] = 3
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)
        self.assertFalse(self.output.exists())

    def test_non_english_output_language_is_rejected_before_writing(self):
        self.bank["language"] = "ar"
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)
        self.assertFalse(self.output.exists())

    def test_difficulty_quotas_are_checked_per_type(self):
        self.bank["mcqs"][0]["difficulty"] = "hard"
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)

    def test_duplicate_stem_is_rejected_across_types(self):
        self.bank["essays"][0]["question"] = "  HOW  MANY sides does a triangle have?  "
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)

    def test_incomplete_answers_and_sources_are_rejected(self):
        cases = [("correct", "E"), ("explanation", ""), ("sources", []),
                 ("answer_verified", False), ("difficulty", "extreme"),
                 ("options", ["same", "same", "other", "last"])]
        for field, replacement in cases:
            with self.subTest(field=field):
                bank = sample_bank()
                bank["mcqs"][0][field] = replacement
                with self.assertRaises(ValueError):
                    formatter.build_docx(bank, self.output)

    def test_false_statement_requires_correction(self):
        del self.bank["true_false"][1]["correction"]
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)

    def test_invalid_true_false_and_empty_essay_are_rejected(self):
        self.bank["true_false"][0]["correct"] = "Maybe"
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)
        self.bank = sample_bank()
        self.bank["essays"][0]["model_answer"] = []
        with self.assertRaises(ValueError):
            formatter.build_docx(self.bank, self.output)

    def test_pdf_mode_rejects_external_sources_and_invalid_page(self):
        for source in [
            {"kind": "web", "title": "Geometry", "url": "https://example.org/geometry", "evidence": "Three sides."},
            {"kind": "pdf", "document": "practice.pdf", "pdf_page": 0, "evidence": "Three sides."},
            {"kind": "pdf", "document": "practice.pdf", "pdf_page": 1, "evidence": ""}
        ]:
            with self.subTest(source=source):
                bank = sample_bank()
                bank["mcqs"][0]["sources"] = [source]
                with self.assertRaises(ValueError):
                    formatter.build_docx(bank, self.output)

    def test_expanded_mode_marks_web_and_general_knowledge(self):
        self.bank["source_mode"] = "pdf_web_general"
        self.bank["mcqs"][0]["sources"] = [
            {"kind": "web", "title": "Geometry", "url": "https://example.org/geometry",
             "evidence": "Three sides."},
            {"kind": "general", "evidence": "General geometric knowledge.",
             "verified_against": "practice.pdf PDF page 1"}
        ]
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        self.assertTrue(document.tables, "The default answer key must be present")
        citations = "\n".join(row.cells[-1].text for row in document.tables[0].rows)
        self.assertIn("https://example.org/geometry", citations)
        self.assertIn("Outside PDF", citations)
        self.assertIn("General knowledge", citations)

    def test_explicit_student_only_request_omits_answer_key(self):
        self.bank["include_answer_key"] = False
        formatter.build_docx(self.bank, self.output)
        self.assertEqual(len(Document(self.output).tables), 0)

    def test_uploaded_slides_are_cited_by_slide_number(self):
        self.bank['source_mode'] = 'source_only'
        source = {'kind': 'slides', 'document': 'geometry.pptx', 'slide': 7,
                  'evidence': 'The slide defines triangle and square side counts.'}
        for questions in (self.bank['mcqs'], self.bank['true_false'], self.bank['essays']):
            for question in questions:
                question['sources'] = [copy.deepcopy(source)]
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        key_text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('geometry.pptx', key_text)
        self.assertIn('slide 7', document.tables[0].rows[1].cells[3].text)
        self.assertNotIn('PDF page', document.tables[0].rows[1].cells[3].text)

    def test_research_document_requires_and_prints_a_section_locator(self):
        self.bank['source_mode'] = 'source_only'
        self.bank['essays'][0]['sources'] = [
            {'kind': 'document_inference', 'document': 'paper.docx',
             'locator': 'Methods, Table 2', 'evidence': 'The methods table supports the comparison.'}
        ]
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        citations = '\n'.join(row.cells[3].text for row in document.tables[0].rows)
        self.assertIn('Methods, Table 2', citations)
        self.assertIn('Inference from source', citations)
        self.bank['essays'][0]['sources'][0]['locator'] = ''
        with self.assertRaises(ValueError):
            formatter.validate_bank(self.bank)

    def test_invalid_slide_indices_and_document_pages_are_rejected(self):
        self.bank['source_mode'] = 'source_only'
        for field, kind in [('slide', 'slides'), ('page', 'document')]:
            for invalid in (0, -1, True, '7'):
                with self.subTest(field=field, invalid=invalid):
                    source = {'kind': kind, 'document': 'source', 'locator': 'Methods',
                              field: invalid, 'evidence': 'Supporting text.'}
                    self.bank['mcqs'][0]['sources'] = [source]
                    with self.assertRaises(ValueError):
                        formatter.validate_bank(self.bank)

    def test_source_only_rejects_external_content(self):
        self.bank['source_mode'] = 'source_only'
        for source in [
            {'kind': 'web', 'title': 'Geometry', 'url': 'https://example.org', 'evidence': 'Three sides.'},
            {'kind': 'general', 'verified_against': 'An external book', 'evidence': 'Three sides.'}
        ]:
            with self.subTest(kind=source['kind']):
                self.bank['mcqs'][0]['sources'] = [source]
                with self.assertRaises(ValueError):
                    formatter.build_docx(self.bank, self.output)
                self.assertFalse(self.output.exists())

    def test_source_web_general_allows_clearly_marked_external_content(self):
        self.bank['source_mode'] = 'source_web_general'
        self.bank['mcqs'][0]['sources'] = [
            {'kind': 'web', 'title': 'Geometry', 'url': 'https://example.org/geometry', 'evidence': 'Three sides.'}
        ]
        formatter.build_docx(self.bank, self.output)
        citations = '\n'.join(row.cells[3].text for row in Document(self.output).tables[0].rows)
        self.assertIn('Outside uploaded material', citations)
        self.assertIn('https://example.org/geometry', citations)

    def test_sources_are_compact_and_identified_only_beside_the_key(self):
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        texts = [p.text for p in document.paragraphs]
        key_index = texts.index('Model Answer')
        self.assertNotIn('practice.pdf', '\n'.join(texts[:key_index]))
        self.assertIn('practice.pdf', '\n'.join(texts[key_index:]))
        self.assertEqual(document.tables[0].rows[1].cells[3].text, 'S1: PDF page 1')

    def test_multiple_uploaded_sources_get_distinct_reference_ids(self):
        self.bank['source_mode'] = 'source_only'
        self.bank['mcqs'][0]['sources'].append(
            {'kind': 'slides_inference', 'document': 'notes.pptx', 'slide': 3, 'evidence': 'Three sides.'})
        formatter.build_docx(self.bank, self.output)
        document = Document(self.output)
        legend = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('S1:', legend)
        self.assertIn('S2:', legend)
        row = next(row for row in document.tables[0].rows if 'slide 3' in row.cells[3].text)
        self.assertIn('PDF page 1', row.cells[3].text)
        self.assertIn('Inference from slides', row.cells[3].text)

    def test_options_keep_template_indentation_without_extra_bullets(self):
        formatter.build_docx(self.bank, self.output)
        options = [p for p in Document(self.output).paragraphs if p.style.name == 'List Bullet 2']
        self.assertEqual(len(options), 8)
        self.assertTrue(all(p._p.pPr.numPr.numId.val == 0 for p in options))

    def test_answer_table_has_readable_fonts_and_a_repeating_header(self):
        formatter.build_docx(self.bank, self.output)
        table = Document(self.output).tables[0]
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
        self.assertTrue(all(run.font.size.pt == 10 for row in table.rows for cell in row.cells
                            for paragraph in cell.paragraphs for run in paragraph.runs))

    def test_difficulty_is_shuffled_without_changing_quotas(self):
        self.bank["mcqs"] = []
        for difficulty in ("easy", "medium", "hard"):
            for index in range(4):
                question = copy.deepcopy(sample_bank()["mcqs"][0])
                question["question"] = f"Practice {difficulty} {index}?"
                question["difficulty"] = difficulty
                self.bank["mcqs"].append(question)
        self.bank["counts"]["mcqs"] = 12
        self.bank["difficulty_counts"]["mcqs"] = {"easy": 4, "medium": 4, "hard": 4}
        original = [q["question"] for q in self.bank["mcqs"]]
        observed = []
        for _ in range(3):
            formatter.build_docx(self.bank, self.output)
            stems = [p.text.split(". ", 1)[-1] for p in Document(self.output).paragraphs
                     if p.style.name == "List Number"][:12]
            observed.append(stems)
            self.assertCountEqual(stems, original)
        self.assertTrue(any(stems != original for stems in observed))


if __name__ == "__main__":
    unittest.main()
